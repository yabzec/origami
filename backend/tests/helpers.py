import zipfile
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont

FONT = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"


def make_text_image(
    path: Path, text: str = "FATTURA 2026", size: tuple[int, int] = (1200, 400)
) -> Path:
    img = Image.new("RGB", size, "white")
    draw = ImageDraw.Draw(img)
    draw.text((60, size[1] // 3), text, fill="black", font=ImageFont.truetype(FONT, 72))
    img.save(path)
    return path


import uuid


def basis_vector(index: int, dim: int = 1536) -> list[float]:
    vector = [0.0] * dim
    vector[index] = 1.0
    return vector


def seed_document(session, title: str, chunk_specs: list[dict], **doc_kwargs):
    """Insert a ready Document plus chunks. Each spec: {content, embedding?, page_number?, source?}."""
    from app.models import Chunk, ChunkSource, DocStatus, DocType, Document

    doc = Document(
        title=title,
        doc_type=doc_kwargs.pop("doc_type", DocType.text),
        status=doc_kwargs.pop("status", DocStatus.ready),
        **doc_kwargs,
    )
    session.add(doc)
    session.commit()
    session.refresh(doc)
    for index, spec in enumerate(chunk_specs):
        session.add(
            Chunk(
                document_id=doc.id,
                chunk_index=index,
                page_number=spec.get("page_number"),
                attachment_id=spec.get("attachment_id"),
                source=spec.get("source", ChunkSource.content),
                content=spec["content"],
                embedding=spec.get("embedding"),
            )
        )
    session.commit()
    return doc


_ODT_CONTENT = """<?xml version="1.0" encoding="UTF-8"?>
<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" office:version="1.2"><office:body><office:text><text:p>{text}</text:p></office:text></office:body></office:document-content>"""

_ODT_MANIFEST = """<?xml version="1.0" encoding="UTF-8"?>
<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" manifest:version="1.2"><manifest:file-entry manifest:full-path="/" manifest:media-type="application/vnd.oasis.opendocument.text"/><manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/></manifest:manifest>"""


def make_docx(path: Path, pages: list[str]) -> Path:
    """A .docx with one paragraph per entry and a hard page break between entries."""
    import docx

    document = docx.Document()
    for index, text in enumerate(pages):
        if index:
            document.add_page_break()
        document.add_paragraph(text)
    document.save(path)
    return path


def make_odt(path: Path, text: str) -> Path:
    """A minimal single-paragraph OpenDocument text file (mimetype entry first, stored)."""
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr(
            zipfile.ZipInfo("mimetype"),
            "application/vnd.oasis.opendocument.text",
            compress_type=zipfile.ZIP_STORED,
        )
        archive.writestr("content.xml", _ODT_CONTENT.format(text=text), compress_type=zipfile.ZIP_DEFLATED)
        archive.writestr("META-INF/manifest.xml", _ODT_MANIFEST, compress_type=zipfile.ZIP_DEFLATED)
    return path


DX_STORAGE = "1.2.840.10008.5.1.4.1.1.1.1"


def make_dicom(
    path: Path,
    *,
    study_uid: str = "1.2.826.0.1.1",
    series_number: int = 1,
    instance_number: int = 1,
    rows: int = 64,
    cols: int = 48,
    photometric: str = "MONOCHROME2",
    window: tuple[float, float] | None = None,
    frames: int = 1,
    color: bool = False,
    rle: bool = False,
    **tags,
) -> Path:
    """A synthetic DX image: a left-to-right 12-bit gradient (or flat red RGB). `Tag=None` removes a tag."""
    import numpy as np
    from pydicom.dataset import FileDataset, FileMetaDataset
    from pydicom.uid import ExplicitVRLittleEndian, RLELossless, generate_uid

    meta = FileMetaDataset()
    meta.MediaStorageSOPClassUID = DX_STORAGE
    meta.MediaStorageSOPInstanceUID = generate_uid()
    meta.TransferSyntaxUID = ExplicitVRLittleEndian
    ds = FileDataset(str(path), {}, file_meta=meta, preamble=b"\0" * 128)
    ds.SOPClassUID = DX_STORAGE
    ds.SOPInstanceUID = meta.MediaStorageSOPInstanceUID
    ds.StudyInstanceUID = study_uid
    ds.SeriesInstanceUID = f"{study_uid}.{series_number}"
    ds.SeriesNumber = series_number
    ds.InstanceNumber = instance_number
    ds.Modality = "DX"
    ds.PatientName = "ROSSI^MARIO"
    ds.PatientBirthDate = "19800101"
    ds.PatientSex = "M"
    ds.PatientID = "P-1"
    ds.AccessionNumber = "A-1"
    ds.StudyDate = "20260828"
    ds.StudyDescription = "Rx gamba (tibia e perone) - 88.27.3 Destro/a"
    ds.SeriesDescription = "AP"
    ds.InstitutionName = "Ospedale Test"
    ds.BodyPartExamined = "TIBIA FIBULA"
    ds.Rows, ds.Columns = rows, cols
    if color:
        ds.SamplesPerPixel, ds.PhotometricInterpretation, ds.PlanarConfiguration = 3, "RGB", 0
        ds.BitsAllocated, ds.BitsStored, ds.HighBit, ds.PixelRepresentation = 8, 8, 7, 0
        frame = np.zeros((rows, cols, 3), dtype=np.uint8)
        frame[..., 0] = 200
    else:
        ds.SamplesPerPixel, ds.PhotometricInterpretation = 1, photometric
        ds.BitsAllocated, ds.BitsStored, ds.HighBit, ds.PixelRepresentation = 16, 12, 11, 0
        frame = np.tile(np.linspace(0, 4095, cols).astype(np.uint16), (rows, 1))
    data = np.stack([frame] * frames) if frames > 1 else frame
    if frames > 1:
        ds.NumberOfFrames = frames
    if window is not None:
        ds.WindowCenter, ds.WindowWidth = window
    for key, value in tags.items():
        if value is None:
            if key in ds:
                delattr(ds, key)
        else:
            setattr(ds, key, value)
    ds.PixelData = data.tobytes()
    if rle:
        ds.compress(RLELossless)
    path.parent.mkdir(parents=True, exist_ok=True)
    ds.save_as(path, enforce_file_format=True)
    return path


def make_text_pdf(path: Path, lines: list[str]) -> Path:
    """A one-page PDF with a real text layer (Helvetica); no lines gives a page with no text."""
    content = "BT /F1 12 Tf 72 720 Td 14 TL " + " ".join(f"({line}) Tj T*" for line in lines) + " ET"
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        "/Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(content)} >>\nstream\n{content}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n{body}\nendobj\n".encode()
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(bytes(out))
    return path
