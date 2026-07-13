import docx as docx_lib

from app.services.extract import (
    extract_docx,
    extract_pdf_text,
    extract_text_file,
    pdf_needs_ocr,
)
from app.services.ocr import ocr_image
from tests.helpers import make_text_image


def test_extract_pdf_text_reads_text_layer(tmp_path):
    img = make_text_image(tmp_path / "p.png", "CONTRATTO AFFITTO")
    pdf_bytes, _ = ocr_image(img, "ita+eng")
    pdf = tmp_path / "doc.pdf"
    pdf.write_bytes(pdf_bytes)

    pages = extract_pdf_text(pdf)
    assert len(pages) == 1
    assert pages[0][0] == 1
    assert "CONTRATTO" in pages[0][1].upper()


def test_pdf_needs_ocr_heuristic():
    assert pdf_needs_ocr([(1, ""), (2, "ab")]) is True
    assert pdf_needs_ocr([(1, "x" * 200)]) is False
    assert pdf_needs_ocr([]) is True


def test_extract_text_file(tmp_path):
    f = tmp_path / "note.md"
    f.write_text("# Titolo\n\nContenuto della nota.", encoding="utf-8")
    assert "Contenuto della nota." in extract_text_file(f)


def test_extract_docx(tmp_path):
    path = tmp_path / "doc.docx"
    d = docx_lib.Document()
    d.add_paragraph("Primo paragrafo.")
    d.add_paragraph("Secondo paragrafo.")
    d.save(path)
    text = extract_docx(path)
    assert "Primo paragrafo." in text
    assert "\n\n" in text
